"""Downloads the RBI files with the updater and keeps only the Andhra Pradesh / Telangana rows of e-STATES (for inspection)."""
import asyncio, csv, sys, os
sys.path.insert(0, "tools")
import update_finance as U
from openpyxl import load_workbook
asyncio.run(U.download_all("rbi-files"))
wb = load_workbook("rbi-files/estates.xlsx", read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
w = csv.writer(open("out/estates_apts.csv", "w", newline=""))
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 3 or (row and row[1] in U.STATES.values()):
        w.writerow(row[:8])
print("ok")
