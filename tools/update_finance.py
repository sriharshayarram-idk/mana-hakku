"""
Mana Hakku · State finances updater.

Finds the latest RBI publications, downloads them with a real (headless) browser,
and rebuilds web/data/finance.json for Andhra Pradesh and Telangana:

  * State Finances: A Study of Budgets -> e-STATES database (Budget Estimates, Revised Estimates
    and Accounts for every budget head and year) and Statements 19/20 (outstanding liabilities).
  * Handbook of Statistics on Indian States -> GSDP (current prices), per capita NSDP, CPI inflation.

All money is in ₹ crore unless noted. Run: python tools/update_finance.py [--from-dir DIR]
"""
import asyncio, datetime, io, json, math, os, re, sys
from openpyxl import load_workbook

OUT = os.path.join(os.path.dirname(__file__), "..", "web", "data", "finance.json")
SF_PAGE = "https://www.rbi.org.in/Scripts/AnnualPublications.aspx?head=State%20Finances%20:%20A%20Study%20of%20Budgets"
HSIS_PAGE = "https://www.rbi.org.in/Scripts/AnnualPublications.aspx?head=Handbook%20of%20Statistics%20on%20Indian%20States"
STATES = {"ap": "Andhra Pradesh", "ts": "Telangana"}
FIRST_YEAR = 2014

# Sector categories: revenue expenditure (Appendix 2) + capital outlay and loans (Appendix 4), by RBI code.
CATS = [
    ("total", ["2|Total", "4|I", "4|IV"]),
    ("dev", ["2|I", "4|I.1.a.1", "4|I.1.a.2", "4|I.1.a.3", "4|I.1.a.4", "4|I.1.a.5", "4|I.1.a.6", "4|I.1.a.7", "4|I.1.a.8",
             "4|I.1.b.1", "4|I.1.b.2", "4|I.1.b.4", "4|I.1.b.5", "4|I.1.b.6", "4|I.1.b.7", "4|I.1.b.8"]),
    ("education", ["2|I.A.1", "4|I.1.a.1", "4|IV.1.a.1"]),
    ("health", ["2|I.A.2", "2|I.A.3", "4|I.1.a.2", "4|I.1.a.3", "4|IV.1.a.2", "4|IV.1.a.3"]),
    ("water", ["2|I.A.4", "4|I.1.a.4", "4|IV.1.a.4"]),
    ("housing", ["2|I.A.5", "4|I.1.a.5", "4|IV.1.a.5"]),
    ("urban", ["2|I.A.6", "4|I.1.a.6"]),
    ("welfare", ["2|I.A.7", "4|I.1.a.7"]),
    ("labour", ["2|I.A.8"]),
    ("socialsec", ["2|I.A.9", "4|I.1.a.8"]),
    ("nutrition", ["2|I.A.10"]),
    ("calamity", ["2|I.A.11"]),
    ("agri", ["2|I.B.1", "4|I.1.b.1", "4|IV.1.b.1", "4|IV.1.b.2", "4|IV.1.b.3", "4|IV.1.b.4"]),
    ("rural", ["2|I.B.2", "4|I.1.b.2", "4|IV.1.b.9"]),
    ("irrigation", ["2|I.B.4", "4|I.1.b.4", "4|IV.1.b.5"]),
    ("energy", ["2|I.B.5", "4|I.1.b.5", "4|IV.1.b.6"]),
    ("industry", ["2|I.B.6", "4|I.1.b.6", "4|IV.1.b.7", "4|IV.1.b.8"]),
    ("transport", ["2|I.B.7", "4|I.1.b.7", "4|I.1.b.8"]),
]


# Income (Appendix 1) and spending by kind (Appendix 2 = day-to-day, Appendix 4 = building assets / loans given).
FLOW_FIELDS = ["ownTax", "centralShare", "ownNonTax", "grants", "revExp", "capOutlay", "loansGiven"]
FLOW_CODES = [(1, "I.A"), (1, "I.B"), (1, "II.C"), (1, "II.D"), (2, "Total"), (4, "I"), (4, "IV")]


def rnd(x, dec=0):
    """Round half away from zero (the RBI tables and the app use this, not banker's rounding)."""
    if x is None:
        return None
    f = 10 ** dec
    v = math.floor(abs(x) * f + 0.5) / f
    v = v if x >= 0 else -v
    return int(v) if dec == 0 else v


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


# ---------------------------------------------------------------- download
async def download_all(dest):
    from playwright.async_api import async_playwright
    found = {}
    async with async_playwright() as p:
        b = await p.chromium.launch()
        ctx = await b.new_context(user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                                             "(KHTML, like Gecko) Chrome/126.0 Safari/537.36")
        pg = await ctx.new_page()

        async def rows_with_links(url):
            await pg.goto(url, wait_until="networkidle", timeout=120000)
            await pg.wait_for_timeout(3000)
            title = await pg.evaluate("()=>{const h=[...document.querySelectorAll('h2,h3,.tableheader,td')].map(e=>e.innerText.trim()).find(t=>/20\\d\\d-\\d\\d/.test(t)&&/(Budgets|Handbook)/i.test(t));return h||document.title}")
            rows = await pg.evaluate("""()=>[...document.querySelectorAll('tr')].map(tr=>({t:tr.innerText.replace(/\\s+/g,' ').trim(),
                       x:[...tr.querySelectorAll('a')].map(a=>a.href).filter(h=>/\\.xlsx$/i.test(h))})).filter(r=>r.x.length)""")
            links = await pg.eval_on_selector_all("a", "els=>els.map(e=>e.href)")
            return title, rows, links

        sf_title, sf_rows, sf_links = await rows_with_links(SF_PAGE)
        hs_title, hs_rows, hs_links = await rows_with_links(HSIS_PAGE)

        def by_title(rows, pattern, fallback=None, links=(), exclude=None):
            for r in rows:
                if re.search(pattern, r["t"], re.I) and not (exclude and re.search(exclude, r["t"], re.I)):
                    return r["x"][0]
            if fallback:
                for l in links:
                    if re.search(fallback, l, re.I):
                        return l
            return None

        want = {
            "estates": next((l for l in sf_links if re.search(r"/ESTATES[^/]*\.xlsx$", l, re.I)), None),
            "liab": by_title(sf_rows, r"Total Outstanding Liabilities of State Governments(?!.*GSDP)", r"/19_ST[^/]*\.xlsx$", sf_links),
            "liabpct": by_title(sf_rows, r"Total Outstanding Liabilities of State Governments.*GSDP", r"/20_ST[^/]*\.xlsx$", sf_links),
            "gsdp": by_title(hs_rows, r"Gross State Domestic Product \(Current Prices\)", None),
            "pc": by_title(hs_rows, r"Per Capita Net State Domestic Product \(Current Prices\)", None),
            "nsdp": by_title(hs_rows, r"Net State Domestic Product \(Current Prices\)", None, exclude=r"Per Capita"),
            "cpi": by_title(hs_rows, r"Average Inflation \(CPI\)\s*[-–]\s*General", None),
        }
        missing = [k for k, v in want.items() if not v]
        if missing:
            raise SystemExit("Could not find links for: " + ", ".join(missing))
        # Pass the site's bot check once on the documents host, then download with the same cookies.
        try:
            await pg.goto("https://rbidocs.rbi.org.in/rdocs/Publications/DOCs/", timeout=60000)
        except Exception:
            pass
        await pg.wait_for_timeout(4000)
        os.makedirs(dest, exist_ok=True)
        for k, u in want.items():
            r = await ctx.request.get(u, timeout=180000)
            body = await r.body()
            if r.status != 200 or body[:2] != b"PK":
                raise SystemExit(f"Download failed for {k}: {u} ({r.status}, {len(body)} bytes)")
            open(os.path.join(dest, k + ".xlsx"), "wb").write(body)
            found[k] = u
        await b.close()
    found["sf_title"], found["hsis_title"] = sf_title, hs_title
    json.dump(found, open(os.path.join(dest, "links.json"), "w"), indent=1)
    return found


# ---------------------------------------------------------------- parse
def parse_estates(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    G = {}
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 7:
            continue
        app, state, head, fy = row[0], row[1], row[2], row[3]
        if state not in STATES.values() or not fy or not isinstance(fy, str) or not re.match(r"^\d{4}-\d{4}$", fy):
            continue
        y = int(fy[:4])
        if y < FIRST_YEAR:
            continue
        a = str(app or "")[-1:]
        code = str(head).split(":")[0].strip()
        k = (state, y, a, code)
        if k not in G:
            G[k] = [num(row[4]), num(row[5]), num(row[6])]  # Accounts, Revised, Budget
    return G


def build_fiscal(G):
    years = sorted({k[1] for k in G})
    def g(s, y, a, c, i):
        v = G.get((s, y, str(a), c))
        return v[i] if v else None
    def ssum(s, y, i, items):
        any_, t = False, 0.0
        for it in items:
            a, c = it[0], it[1]
            sg = it[2] if len(it) > 2 else 1
            v = g(s, y, a, c, i)
            if v is not None:
                any_, t = True, t + sg * v
        return rnd(t) if any_ else None
    def interest(s, y, i):
        t = g(s, y, 2, "II.C.2", i)
        if t is None:
            return None
        comp = sum((g(s, y, 2, c, i) or 0) for c in ["II.C.2.i", "II.C.2.ii", "II.C.2.iii", "II.C.2.iv"])
        if abs(t - comp) <= max(5, t * 0.01):
            return rnd(t)
        alt = g(s, y, 2, "II.C.1", i)  # rows shifted by one in the source for a few cells
        return rnd(alt) if alt is not None else rnd(comp)
    def cat_val(s, y, i, codes):
        any_, t = False, 0.0
        for kc in codes:
            a, c = kc.split("|")
            v = g(s, y, a, c, i)
            if v is not None:
                any_, t = True, t + v
        return rnd(t) if any_ else None

    out = {"years": [f"{y}-{str(y + 1)[2:]}" for y in years]}
    for key, name in STATES.items():
        yr, cats, flow = [], [], []
        for y in years:
            row = []
            for i in (2, 0):  # Budget Estimates, Accounts
                te = ssum(name, y, i, [(2, "Total"), (4, "I"), (4, "IV")])
                rr = ssum(name, y, i, [(1, "Total")])
                nd = ssum(name, y, i, [(3, "III"), (3, "XI")]) or 0
                row.append([rr, interest(name, y, i),
                            ssum(name, y, i, [(3, "I"), (3, "I.6", -1), (3, "II"), (3, "II.5", -1)]),
                            ssum(name, y, i, [(4, "II"), (4, "II.6", -1), (4, "III"), (4, "III.5", -1)]),
                            (te - rr - nd) if te is not None and rr is not None else None,
                            ssum(name, y, i, [(2, "II.E")])])
            yr.append(row)
            cats.append([[cat_val(name, y, 2, codes), cat_val(name, y, 0, codes)] for _, codes in CATS])
            # Money in and money out, family-budget style (FLOW_FIELDS order), [Budget Estimates, Accounts]
            flow.append([[ssum(name, y, i, [(a, c)]) for a, c in FLOW_CODES] for i in (2, 0)])
        out[key] = {"yr": yr, "cat": cats, "flow": flow}
    return out


def state_row(rows, name, col=1):
    rx = re.compile(r"^(\d+\.\s*)?" + re.escape(name) + r"\s*$", re.I)
    for r in rows:
        if r and len(r) > col and r[col] is not None and rx.match(str(r[col]).strip()):
            return r
    return None


def parse_hsis(path, years, scale=1.0, dec=0):
    wb = load_workbook(path, read_only=True, data_only=True)
    res = {k: {} for k in STATES}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = next((r for r in rows if r and any(re.match(r"^20\d\d-\d\d$", str(v or "").strip()) for v in r)), None)
        if not hdr:
            continue
        for k, name in STATES.items():
            r = state_row(rows, name)
            if not r:
                continue
            for i, v in enumerate(hdr):
                yv = str(v or "").strip()
                if re.match(r"^20\d\d-\d\d$", yv) and i < len(r):
                    n = num(r[i])
                    res[k][yv] = None if n is None else rnd(n / scale, dec)
    return {k: [res[k].get(y) for y in years] for k in STATES}


def parse_liab(path, dec=0):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    hdr = next(r for r in rows if r and any(str(v or "").strip() == "2008" for v in r))
    cols = [(i, str(v).strip()) for i, v in enumerate(hdr) if v is not None and re.match(r"^20\d\d", str(v).strip()) and int(str(v).strip()[:4]) >= FIRST_YEAR + 1]
    labels = ["Mar " + c[1].replace("(RE)", "(revised)").replace("(BE)", "(budget)") for c in cols]
    out = {}
    for k, name in STATES.items():
        r = state_row(rows, name)
        out[k] = [None if num(r[i]) is None else rnd(num(r[i]), dec) for i, _ in cols]
    return labels, out


def build(dir_, links):
    G = parse_estates(os.path.join(dir_, "estates.xlsx"))
    f = build_fiscal(G)
    years = f["years"]
    gsdp = parse_hsis(os.path.join(dir_, "gsdp.xlsx"), years, scale=100, dec=0)   # ₹ lakh -> ₹ crore
    pc = parse_hsis(os.path.join(dir_, "pc.xlsx"), years, 1, 0)
    nsdp = parse_hsis(os.path.join(dir_, "nsdp.xlsx"), years, scale=100, dec=0)  # ₹ lakh -> ₹ crore
    # population (lakh) = NSDP / per-capita NSDP
    pop = {k: [rnd(n * 1e7 / p / 1e5, 1) if n and p else None for n, p in zip(nsdp[k], pc[k])] for k in STATES}
    cpi = parse_hsis(os.path.join(dir_, "cpi.xlsx"), years, 1, 1)
    lyears, liab = parse_liab(os.path.join(dir_, "liab.xlsx"), 0)
    _, liabpct = parse_liab(os.path.join(dir_, "liabpct.xlsx"), 1)
    data = {
        "meta": {
            "updated": datetime.date.today().isoformat(),
            "units": "₹ crore unless noted",
            "sources": {
                "sf": {"title": links.get("sf_title", "State Finances: A Study of Budgets"), "page": SF_PAGE,
                       "files": {k: links[k] for k in ("estates", "liab", "liabpct") if k in links}},
                "hsis": {"title": links.get("hsis_title", "Handbook of Statistics on Indian States"), "page": HSIS_PAGE,
                         "files": {k: links[k] for k in ("gsdp", "nsdp", "pc", "cpi") if k in links}},
            },
            "cats": [c for c, _ in CATS],
            "yrFields": ["revenue", "interest", "borrowed", "repaid", "deficit", "pensions"],
            "flowFields": FLOW_FIELDS,
        },
        "years": years, "liabYears": lyears,
    }
    for k in STATES:
        data[k] = {"gsdp": gsdp[k], "nsdp": nsdp[k], "pop": pop[k], "pc": pc[k], "cpi": cpi[k], "liab": liab[k], "liabPct": liabpct[k],
                   "yr": f[k]["yr"], "cat": f[k]["cat"], "flow": f[k]["flow"]}
    return data


def sane(d):
    """Refuse to publish obviously broken data."""
    for k in STATES:
        s = d[k]
        assert len(d["years"]) >= 10, "too few years"
        acc = [y[1][0] for y in s["yr"] if y[1][0]]
        assert len(acc) >= 8, f"{k}: too few accounts years"
        assert sum(1 for v in s["gsdp"] if v) >= 8, f"{k}: GSDP missing"
        assert sum(1 for v in s["liab"] if v) >= 8, f"{k}: liabilities missing"
        for y, fl in zip(s["yr"], s["flow"]):
            for i in (0, 1):
                if y[i][0] and all(v is not None for v in fl[i][:4]):
                    assert abs(sum(fl[i][:4]) - y[i][0]) <= 5, f"{k}: income parts do not add up"
        tot = [c[0][0] for c in s["cat"] if c[0][0]]
        assert tot and min(tot) > 50000, f"{k}: totals look wrong"
    return True


def main():
    args = sys.argv[1:]
    dir_ = "rbi-files"
    if "--from-dir" in args:
        dir_ = args[args.index("--from-dir") + 1]
        links = json.load(open(os.path.join(dir_, "links.json"))) if os.path.exists(os.path.join(dir_, "links.json")) else {}
    else:
        links = asyncio.run(download_all(dir_))
    data = build(dir_, links)
    sane(data)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    old = None
    if os.path.exists(OUT):
        old = json.load(open(OUT))
    same = old is not None and {k: v for k, v in old.items() if k != "meta"} == {k: v for k, v in data.items() if k != "meta"}
    if same:
        data = old  # keep the original "updated" date
    txt = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    if not same:
        open(OUT, "w").write(txt)
    # keep the copy built into the app page in step (no change if already equal)
    page = os.path.join(os.path.dirname(__file__), "..", "app", "src", "main", "assets", "index.html")
    h = open(page, encoding="utf-8").read()
    a, b = "/*FIN-DATA*/", "/*END-FIN-DATA*/"
    if a in h and b in h:
        h = h[:h.index(a) + len(a)] + txt + h[h.index(b):]
        open(page, "w", encoding="utf-8").write(h)
    print("unchanged" if same else "UPDATED", OUT)


if __name__ == "__main__":
    main()
